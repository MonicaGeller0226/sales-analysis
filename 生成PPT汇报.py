# -*- coding: utf-8 -*-
"""
生成商务风销售分析PPT（≤10页），嵌入已生成的图表PNG
输出: D:\Python\Python_item\PythonProject\销售数据分析\PPT\7月销售数据分析_汇报.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

BASE = r'D:\Python\Python_item\PythonProject\销售数据分析'
CHART = os.path.join(BASE, '图表')
OUT_DIR = os.path.join(BASE, 'PPT')
OUT = os.path.join(OUT_DIR, '7月销售数据分析_汇报.pptx')
os.makedirs(OUT_DIR, exist_ok=True)

# ---------- 配色（商务风） ----------
NAVY   = RGBColor(0x1F, 0x4E, 0x79)   # 深蓝 主色
BLUE   = RGBColor(0x2E, 0x74, 0xB5)   # 中蓝
LTBLUE = RGBColor(0xD6, 0xE4, 0xF0)   # 浅蓝底
RED    = RGBColor(0xC0, 0x00, 0x00)   # 红（强调/问题）
GOLD   = RGBColor(0xBF, 0x8F, 0x00)   # 金（建议/亮点）
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x33, 0x33, 0x33)
GRAY   = RGBColor(0x77, 0x77, 0x77)

FONT = '微软雅黑'
SW, SH = Inches(13.333), Inches(7.5)  # 16:9

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]

def set_font(run, size=18, bold=False, color=DARK, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    # 中文字体
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', name)

def add_slide():
    return prs.slides.add_slide(BLANK)

def add_rect(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp

def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: list of (text, size, bold, color) or list of list(paragraph)"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    first = True
    for para in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        items = para if isinstance(para, list) else [para]
        for (t, s, b, c) in items:
            run = p.add_run(); run.text = t
            set_font(run, size=s, bold=b, color=c)
    return tb

def header(slide, title, subtitle=None):
    """顶部通栏标题"""
    add_rect(slide, 0, 0, SW, Inches(0.9), NAVY)
    add_rect(slide, 0, Inches(0.9), SW, Inches(0.06), GOLD)
    tb = add_text(slide, Inches(0.5), Inches(0.12), Inches(9), Inches(0.7),
                  [[(title, 26, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, Inches(9.5), Inches(0.12), Inches(3.4), Inches(0.7),
                 [[(subtitle, 13, False, LTBLUE)]], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

def footer(slide, page):
    add_text(slide, Inches(0.4), Inches(7.1), Inches(6), Inches(0.35),
             [[("理想汽车 · 2026年7月销售分析", 10, False, GRAY)]])
    add_text(slide, Inches(12.5), Inches(7.1), Inches(0.7), Inches(0.35),
             [[(str(page), 12, False, GRAY)]], align=PP_ALIGN.RIGHT)

def add_picture_fit(slide, img, cx, cy, max_w, max_h):
    """居中等比缩放放入图片"""
    from PIL import Image
    iw, ih = Image.open(img).size
    scale = min(max_w/iw, max_h/ih)
    w = int(iw*scale); h = int(ih*scale)
    left = int(cx - w/2); top = int(cy - h/2)
    slide.shapes.add_picture(img, left, top, w, h)

def bullet_list(slide, x, y, w, h, items, size=15, gap=6, marker='• ', mcolor=NAVY):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        if isinstance(it, tuple):   # (text, color)
            t, c = it
        else:
            t, c = it, DARK
        r1 = p.add_run(); r1.text = marker; set_font(r1, size=size, bold=True, color=mcolor)
        r2 = p.add_run(); r2.text = t;      set_font(r2, size=size, bold=False, color=c)
    return tb

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

# ============================================================
# 1. 封面
# ============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, NAVY)
add_rect(s, 0, Inches(4.9), SW, Inches(0.08), GOLD)
add_text(s, Inches(1), Inches(2.3), Inches(11), Inches(1.2),
         [[("理想汽车 · 2026年7月销售数据分析", 44, True, WHITE)]])
add_text(s, Inches(1), Inches(3.6), Inches(11), Inches(0.6),
         [[("营销工作总结 · 问题诊断 · 改进建议", 22, False, LTBLUE)]])
add_text(s, Inches(1), Inches(5.2), Inches(11), Inches(0.5),
         [[("汇报对象：销售主管 & 销售人员        |        汇报周期：2026-07-01 至 07-28", 15, False, WHITE)]])
add_text(s, Inches(1), Inches(6.6), Inches(11), Inches(0.5),
         [[("数据来源：7月门店销售明细（553条记录）", 12, False, RGBColor(0xAF,0xC4,0xD9))]])
notes(s, "开场：7月整体销售动能持续向上，本次汇报聚焦数据表现、5大问题诊断与改进方向，供主管与销售团队对齐。")

# ============================================================
# 2. 总体概览（数据卡）
# ============================================================
s = add_slide()
header(s, "月度总体概览", "总览")
# 四个大数字卡
cards = [
    ("490", "总成交台数", "#4C72B0"),
    ("3,692", "总线索量", "#DD8452"),
    ("11,263", "总到店人数", "#55A868"),
    ("13.27%", "线索转化率", "#1F4E79"),
]
cw, ch, gap = Inches(2.9), Inches(1.9), Inches(0.25)
x0 = Inches(0.5); y0 = Inches(1.5)
for i, (num, lab, col) in enumerate(cards):
    x = x0 + i*(cw+gap)
    add_rect(s, x, y0, cw, ch, LTBLUE)
    add_rect(s, x, y0, Inches(0.12), ch, BLUE)
    add_text(s, x+Inches(0.3), y0+Inches(0.3), cw-Inches(0.5), Inches(0.9),
             [[(num, 40, True, NAVY)]])
    add_text(s, x+Inches(0.3), y0+Inches(1.3), cw-Inches(0.5), Inches(0.5),
             [[(lab, 14, False, DARK)]])
# 到店转化、关键结论
add_text(s, Inches(0.5), Inches(3.7), Inches(12.3), Inches(0.5),
         [[("整体转化效率：到店转化率 4.35%  ·  线索转化率 13.27%  ·  到店线索率 305%（到店含自然客流）",
            16, True, NAVY)]])
bullet_list(s, Inches(0.5), Inches(4.4), Inches(12.3), Inches(2.4), [
    ("周度销售持续增长：成交台数由 W1 的 91 台升至 W4 的 130 台（+43%），动能向上", DARK),
    ("主力车型 L7/L8 合计占比 87.2%，是绝对走量引擎", DARK),
    ("销售顾问业绩分化显著：王浩 / 张磊两人贡献全司 35.5% 销量", DARK),
    ("到店→成交转化仅 4.35%，是当前最大的销量增长潜力点", RED),
], size=15, gap=8)
footer(s, 2)
notes(s, "先给结论：总额亮眼、周度上扬，但到店转化和顾问分化是隐患。")

# ============================================================
# 3. 时间维度：每日趋势 + 分周对比
# ============================================================
s = add_slide()
header(s, "时间维度：销售趋势与周度表现", "趋势")
add_picture_fit(s, os.path.join(CHART, '1_每日销售趋势.png'),
                int(Inches(6.7)), int(Inches(3.3)), int(Inches(6.2)), int(Inches(3.6)))
add_picture_fit(s, os.path.join(CHART, '2_分周汇总.png'),
                int(Inches(11.7)), int(Inches(3.3)), int(Inches(2.8)), int(Inches(3.6)))
add_text(s, Inches(0.5), Inches(5.6), Inches(12.3), Inches(1.4), [
    [("解读：", 14, True, NAVY), ("成交/到店/线索三线同步上行，节奏稳定；分周对比显示 W1 91 台 → W4 130 台，环比 +43%，销售动能持续走强（W5 为不完整周，仅供参考）。", 14, False, DARK)],
])
footer(s, 3)
notes(s, "强调周度增长 43%；W5 数据截止7-28属不完整周，不纳入趋势结论。")

# ============================================================
# 4. 门店/城市分析
# ============================================================
s = add_slide()
header(s, "门店 / 城市分析", "门店")
add_picture_fit(s, os.path.join(CHART, '3_各城市成交对比.png'),
                int(Inches(4.5)), int(Inches(3.1)), int(Inches(4.4)), int(Inches(3.9)))
add_picture_fit(s, os.path.join(CHART, '11_城市车型热力图.png'),
                int(Inches(10.8)), int(Inches(3.1)), int(Inches(4.7)), int(Inches(3.9)))
bullet_list(s, Inches(0.5), Inches(5.7), Inches(12.3), Inches(1.4), [
    ("上海（133台）为量质双优标杆，成都、重庆转化率靠前但线索量偏少，存在放量空间", DARK),
    ("广州线索量不少（529）但转化率仅 9.3% 全场垫底，属高线索低转化，需重点治理", RED),
])
footer(s, 4)
notes(s, "门店层面点名上海标杆、广州痛点；热力图补充车型分布。")

# ============================================================
# 5. 车型分析
# ============================================================
s = add_slide()
header(s, "车型分析", "车型")
add_picture_fit(s, os.path.join(CHART, '5_各车型销量占比.png'),
                int(Inches(4.3)), int(Inches(3.2)), int(Inches(4.2)), int(Inches(4.1)))
add_picture_fit(s, os.path.join(CHART, '6_各车型转化率.png'),
                int(Inches(10.6)), int(Inches(3.2)), int(Inches(5.0)), int(Inches(4.1)))
add_text(s, Inches(0.5), Inches(6.1), Inches(12.3), Inches(1.2), [
    [("解读：", 14, True, NAVY), ("L7（49.2%）+ L8（38.0%）占 87.2%；L9 高端靠头部城市消化；L6 疲软；MEGA 虽有 178 条线索、576 人到店，但全月 0 成交，为最大疑点。", 14, False, DARK)],
])
footer(s, 5)
notes(s, "重点抛 MEGA 问题：有线索有到店却零成交，需单独复盘。")

# ============================================================
# 6. 销售顾问分析
# ============================================================
s = add_slide()
header(s, "销售顾问分析", "人员")
add_picture_fit(s, os.path.join(CHART, '7_销售顾问业绩排行.png'),
                int(Inches(4.7)), int(Inches(3.2)), int(Inches(4.5)), int(Inches(4.1)))
add_picture_fit(s, os.path.join(CHART, '8_销售顾问产能转化散点.png'),
                int(Inches(10.4)), int(Inches(3.2)), int(Inches(4.9)), int(Inches(4.1)))
bullet_list(s, Inches(0.5), Inches(6.0), Inches(12.3), Inches(1.3), [
    ("王浩（91台）/ 张磊（83台）合计占 35.5%，量、质双优，可作标杆复制", GRAY),
    ("吴敏、郑琪整月 0 成交，李娜转化率仅 7.9%，需专项帮扶", RED),
])
footer(s, 6)
notes(s, "扬长：复制销冠方法论；短板：帮扶零/低转化顾问。")

# ============================================================
# 7. 转化漏斗 + 星期规律
# ============================================================
s = add_slide()
header(s, "转化漏斗与客流规律", "漏斗")
add_picture_fit(s, os.path.join(CHART, '9_整体转化漏斗.png'),
                int(Inches(4.6)), int(Inches(3.4)), int(Inches(4.0)), int(Inches(4.4)))
add_picture_fit(s, os.path.join(CHART, '10_星期几成交规律.png'),
                int(Inches(10.2)), int(Inches(3.4)), int(Inches(5.0)), int(Inches(4.4)))
add_text(s, Inches(0.5), Inches(6.2), Inches(12.3), Inches(1.0), [
    [("关键瓶颈：", 14, True, RED), ("到店 11,263 人 → 成交 490 台，到店转化仅 4.35%。到店客流池极大（且大于线索量，含自然客流），提升到店转化是最有性价比的杠杆。", 14, False, DARK)],
])
footer(s, 7)
notes(s, "漏斗是核心页：到店转化 4.35%，每提升1pp≈多成交113台。")

# ============================================================
# 8. 问题点汇总
# ============================================================
s = add_slide()
header(s, "问题点诊断（Top 5）", "问题")
probs = [
    ("问题 1", "到店转化率偏低", "仅 4.35%，每 100 名到店客户仅约 4 人成交，客流未充分变现"),
    ("问题 2", "广州门店高线索低转化", "线索 529 条不少，但转化率 9.3% 全场垫底，线索质量/跟进存疑"),
    ("问题 3", "MEGA 车型全线 0 成交", "178 条线索、576 人到店却零转化，市场接受度或成交路径存障"),
    ("问题 4", "顾问业绩两极分化", "王浩+张磊占 35.5%；吴敏、郑琪 0 成交，李娜转化仅 7.9%"),
    ("问题 5", "部分车型疲软", "L6 仅 14 台且转化率低，边缘车型贡献有限"),
]
y = Inches(1.35)
for tag, t, desc in probs:
    add_rect(s, Inches(0.5), y, Inches(1.35), Inches(0.95), RED)
    add_text(s, Inches(0.5), y, Inches(1.35), Inches(0.95),
             [[(tag, 14, True, WHITE)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Inches(1.85), y, Inches(11.0), Inches(0.95), LTBLUE)
    add_text(s, Inches(2.05), y+Inches(0.08), Inches(10.6), Inches(0.85),
             [[(t + "：", 14, True, NAVY), (desc, 13, False, DARK)]])
    y += Inches(1.08)
footer(s, 8)
notes(s, "逐条讲5大问题，强调到店转化和广州/顾问/MEGA三项。")

# ============================================================
# 9. 改进建议
# ============================================================
s = add_slide()
header(s, "改进建议（对策）", "行动")
sug = [
    ("提升到店转化", "优化试驾体验与逼单环节；按客流时段排班；建立到店未成交回访机制", GOLD),
    ("广州门店专项整改", "复盘线索质量、规范跟进 SOP；对低效线索重新清洗与分配", GOLD),
    ("MEGA 独立复盘", "单独跟踪线索来源与客户反馈，评估定位与推广策略", GOLD),
    ("顾问帮扶与复制", "沉淀销冠话术/跟进方法论；对 0 转化顾问开展专项培训与带教", GOLD),
    ("放量高转化城市", "向成都、重庆等转化率高但线索少的城市加大投放", GOLD),
    ("主抓走量车型 L7/L8", "重点补给主力车型线索，同时抢救 L6/L9 转化短板", GOLD),
]
y = Inches(1.35)
for t, d, col in sug:
    add_rect(s, Inches(0.5), y, Inches(0.15), Inches(0.95), GOLD)
    add_rect(s, Inches(0.65), y, Inches(11.85), Inches(0.95), WHITE, line=LTBLUE)
    add_text(s, Inches(0.9), y+Inches(0.06), Inches(11.4), Inches(0.85),
             [[(t + "　", 14, True, NAVY), ("—— " + d, 12.5, False, DARK)]])
    y += Inches(1.0)
footer(s, 9)
notes(s, "每条对策对应一个问题的落地方案，主管可据此分配责任人。")

# ============================================================
# 10. 总结/下一步
# ============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, NAVY)
add_rect(s, 0, Inches(1.1), Inches(13.333), Inches(0.07), GOLD)
add_text(s, Inches(1), Inches(1.6), Inches(11), Inches(1.0),
         [[("下一步行动 & 一句话总结", 34, True, WHITE)]])
steps = [
    ("①", "主抓到店转化：目标到店转化率由 4.35% 提升 1 个百分点（≈多成交 113 台）"),
    ("②", "广州门店专项：两周内完成线索质量与跟进 SOP 复盘"),
    ("③", "MEGA 独立诊断：输出专项报告，评估车型定位"),
    ("④", "顾问帮扶：销冠方法论沉淀 + 0/低转化顾问带教计划"),
    ("⑤", "放量高转化城市（成都/重庆），补给 L7/L8 主力线索"),
]
y = Inches(2.9)
for num, t in steps:
    add_text(s, Inches(1), y, Inches(11.5), Inches(0.6),
             [[(num + "  ", 18, True, GOLD), (t, 17, False, WHITE)]])
    y += Inches(0.75)
add_text(s, Inches(1), Inches(6.7), Inches(11), Inches(0.6),
         [[("数据驱动 · 聚焦问题 · 快速行动", 15, True, LTBLUE)]])
notes(s, "收尾：明确责任人分工与优先级，主抓到店转化这个最大杠杆。")

prs.save(OUT)
print('PPT 已生成:', OUT)
print('总页数:', len(prs.slides.__iter__.__self__._sldIdLst))
